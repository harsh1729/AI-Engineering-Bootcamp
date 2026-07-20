from pathlib import Path

from openpyxl import load_workbook

from app.services.document_exceptions import DocumentParsingError
from app.services.parsers.base import BaseDocumentParser


class XlsxParser(BaseDocumentParser):
    def parse(self, file_path: Path) -> str:
        try:
            workbook = load_workbook(str(file_path), read_only=True, data_only=True)
            sections: list[str] = []

            for sheet in workbook.worksheets:
                rows: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                    if cells:
                        rows.append("\t".join(cells))
                if rows:
                    sections.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))

            workbook.close()
            return "\n\n".join(sections).strip()
        except DocumentParsingError:
            raise
        except Exception as exc:
            raise DocumentParsingError(
                f"Failed to parse XLSX '{file_path.name}': {exc}"
            ) from exc
