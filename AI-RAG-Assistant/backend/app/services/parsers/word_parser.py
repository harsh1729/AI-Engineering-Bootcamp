from pathlib import Path

from docx import Document

from app.services.document_exceptions import DocumentParsingError
from app.services.parsers.base import BaseDocumentParser


class DocxParser(BaseDocumentParser):
    def parse(self, file_path: Path) -> str:
        try:
            document = Document(str(file_path))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

            # Tables are a common source of content that paragraph iteration alone misses.
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append("\t".join(cells))

            return "\n".join(paragraphs).strip()
        except DocumentParsingError:
            raise
        except Exception as exc:
            raise DocumentParsingError(
                f"Failed to parse DOCX '{file_path.name}': {exc}"
            ) from exc
