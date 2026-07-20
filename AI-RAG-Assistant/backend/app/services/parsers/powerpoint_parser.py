from pathlib import Path

from pptx import Presentation

from app.services.document_exceptions import DocumentParsingError
from app.services.parsers.base import BaseDocumentParser


class PptxParser(BaseDocumentParser):
    def parse(self, file_path: Path) -> str:
        try:
            presentation = Presentation(str(file_path))
            slides: list[str] = []

            for index, slide in enumerate(presentation.slides, start=1):
                texts: list[str] = []
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue
                    for paragraph in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in paragraph.runs).strip()
                        if line:
                            texts.append(line)
                if texts:
                    slides.append(f"[Slide {index}]\n" + "\n".join(texts))

            return "\n\n".join(slides).strip()
        except DocumentParsingError:
            raise
        except Exception as exc:
            raise DocumentParsingError(
                f"Failed to parse PPTX '{file_path.name}': {exc}"
            ) from exc
